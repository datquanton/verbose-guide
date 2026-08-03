# -*- coding: utf-8 -*-
"""Cross-check the FPT slides against the master model (uploaded, recalculated,
cost of debt 9.0%) and the stock-pick workbook. Fails loudly."""
import openpyxl, zipfile
from lxml import etree
from pptx import Presentation

DECK='/home/user/verbose-guide/MASVN_RS_WM_2H26_outlook_Equity_VN_2026_STBFPT_31July2026.pptx'
PICK='/home/user/verbose-guide/Stock_Pick_and_Forecast_Aug26_MAS_RS_EN_updated.xlsx'
MODEL='/root/.claude/uploads/041665b7-4ff3-507f-a1e8-7a7ed4160156/2a828447-FPT_2Q26.xlsx'
NS='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
TP, PRICE, SHARES, FY25 = 87950., 71700., 1.703507121, 9464.16
fails=[]
def check(n,c,d=''):
    print('%-4s %-56s %s'%('OK' if c else 'FAIL',n,d))
    if not c: fails.append(n)

# --- the model is the source of truth
z=zipfile.ZipFile(MODEL)
def mcell(sheet,ref):
    root=etree.fromstring(z.read('xl/worksheets/%s.xml'%sheet))
    for c in root.iter(NS+'c'):
        if c.get('r')==ref:
            v=c.find(NS+'v'); return float(v.text) if v is not None else None
M={k:mcell('sheet13','%s%d'%(c,r)) for k,(c,r) in
   {'rev26':('H',3),'rev27':('I',3),'rev28':('J',3),'op26':('H',7),'op27':('I',7),'op28':('J',7),
    'pbt26':('H',15),'pbt27':('I',15),'pbt28':('J',15),
    'np26':('H',20),'np27':('I',20),'np28':('J',20)}.items()}
M['tp']=mcell('sheet5','J21'); M['price']=mcell('sheet5','J23'); M['wacc']=mcell('sheet5','C27')
M['kd']=mcell('sheet5','C18'); M['g']=mcell('sheet5','C7'); M['sh']=mcell('sheet5','J19')

prs=Presentation(DECK); num=lambda s: float(s.replace(',',''))
def grab(i):
    sh={x.shape_id:x for x in prs.slides[i].shapes}
    box={sh[12].table.cell(r,0).text: sh[12].table.cell(r,1).text.strip() for r in range(5)}
    rate={sh[13].table.cell(r,0).text: sh[13].table.cell(r,1).text.strip() for r in range(4)}
    t=sh[16].table
    tbl={t.cell(r,0).text:[t.cell(r,c).text.strip() for c in range(4,7)] for r in range(1,12)}
    return box,rate,tbl,sh[15].text_frame.text
en_box,en_rate,en_tbl,en_txt=grab(2)
vn_box,vn_rate,vn_tbl,vn_txt=grab(3)
row=lambda tbl,k: [num(x) for x in next(v for kk,v in tbl.items() if kk.startswith(k))]

# --- slide vs model
for key,lab in [('Revenue','rev'),('Operating profit','op'),('PBT','pbt'),('NPATMI','np')]:
    got=row(en_tbl,key); want=[M['%s%d'%(lab,y)] for y in (26,27,28)]
    check('slide %s row ties to the model'%key, all(abs(a-b)<1 for a,b in zip(got,want)),
          '%s vs %s'%([round(x) for x in got],[round(x) for x in want]))
npat=row(en_tbl,'NPATMI'); eps=row(en_tbl,'EPS'); pe=row(en_tbl,'P/E')
check('shares: model J19 = 1,703,507,121', abs(M['sh']-1703507121)<1)
for i,y in enumerate((26,27,28)):
    check('FY%dF EPS = NPATMI / 1.7035bn shares'%y, abs(npat[i]/SHARES-eps[i])<1.5,
          '%.0f vs %.0f'%(npat[i]/SHARES,eps[i]))
    check('FY%dF P/E = target / EPS'%y, abs(TP/eps[i]-pe[i])<0.06, '%.2f vs %.2f'%(TP/eps[i],pe[i]))
prev=FY25
for i,(y,g) in enumerate(zip((26,27,28),(15.6,15.8,16.6))):
    got=npat[i]/prev*100-100
    check('FY%dF growth %.1f%%'%(y,g), abs(got-g)<0.06, '%.2f%%'%got); prev=npat[i]
# the master model taxes the whole PBT at 15% and carries MI at ~1.49% of PAT
pbt=row(en_tbl,'PBT')
MI_R=(0.0149063,0.0163969,0.0163971)          # the model's MI hack differs by year
for i,y in enumerate((26,27,28)):
    got=pbt[i]*0.85*(1-MI_R[i])
    check('FY%dF PBT -> NPATMI at the model\'s 15%% tax and MI'%y, abs(got-npat[i])<6,
          '%.0f vs %.0f'%(got,npat[i]))
# --- target price and rating box
check('target price 87,950 (model %.0f, rounded to 50)'%M['tp'],
      en_rate['Target price \n(VND, 12M)']=='87,950' and abs(TP-M['tp'])<50)
check('current price = the model\'s', num(en_rate['Current price \n(31/07/26)'])==M['price'],
      '%.0f'%M['price'])
check('expected return = TP/price - 1', abs(TP/PRICE-1-0.227)<0.002, '%.1f%%'%(TP/PRICE*100-100))
check('market cap = price x shares', abs(num(en_box['Market cap (VNDbn)'])-PRICE*SHARES)<2,
      '%.0f'%(PRICE*SHARES))
check('box NPATMI = table FY26F', num(en_box['NPATMI (26F, VNDbn)'])==npat[0])
check('box P/E = table FY26F', num(en_box['P/E (26F, x)'])==pe[0])
check('box EPS growth = NPATMI growth (same share count)',
      abs(num(en_box['EPS Growth (26F, %)'])-15.6)<0.05)
# --- EN and VN identical
for a,b in [('Revenue','Doanh thu'),('PBT','LNTT'),('NPATMI','LNST-CĐTS'),('EPS','EPS'),
            ('P/E','P/E'),('P/B','P/B'),('BVPS','Giá trị sổ sách')]:
    check('EN and VN agree on %s'%a, row(en_tbl,a)==row(vn_tbl,b))
check('EN and VN target price agree',
      en_rate['Target price \n(VND, 12M)']==vn_rate['Giá mục tiêu \n(VND, 12T)'])
# --- narrative
for s,t in [('10,944','FY26F NPATMI'),('12,674','FY27F'),('14,774','FY28F'),('87,950','target'),
            ('71,700','price'),('9.0%','cost of debt'),('11.4%','WACC'),('2,608','associates FY26F'),
            ('+16.7%','Global IT FY27F'),('13.7x','P/E at target'),('26,338','1H26 signings')]:
    check('narrative states %s (%s)'%(s,t), s in en_txt)
# '78,750' and '11.5%' remain, deliberately, as the prior target and prior cost of debt
for old in ['62,900','10,848','12,497','14,459','(+14.6%','(+15.2%','(+15.7%','13.1x',
            '+15.5% in FY27F','+16.0% in FY28F']:
    check('stale figure "%s" gone'%old, old not in en_txt and old not in vn_txt
          and old not in str(en_tbl)+str(vn_tbl)+str(en_rate)+str(vn_rate))
# --- workbook
wb=openpyxl.load_workbook(PICK); ws=wb['Stock Pick']; tp=wb['Target Price and Forecast']
check('workbook NPATMI 26F/27F = deck', ws['F7'].value==npat[0] and ws['G7'].value==npat[1])
check('workbook growth 26F/27F', abs(ws['H7'].value*100-15.63)<0.02 and abs(ws['I7'].value*100-15.81)<0.02)
check('workbook P/E = deck', abs(ws['J7'].value-pe[0])<0.06 and abs(ws['K7'].value-pe[1])<0.06)
check('TP sheet target price', tp['C16'].value==87950)
check('TP sheet NPATMI', tp['D16'].value==npat[0] and tp['E16'].value==npat[1])
print('\n'+'='*76)
print('%d FAILED: %s'%(len(fails),fails) if fails else 'ALL CHECKS PASSED')
