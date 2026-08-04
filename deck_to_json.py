# -*- coding: utf-8 -*-
"""Dump the deck's four slides to JSON so the Word note is built from the same
source as the slides, not retyped."""
import json
from pptx import Presentation

DECK = ('/home/user/verbose-guide/'
        'MASVN_RS_WM_2H26_outlook_Equity_VN_2026_STBFPT_August2026.pptx')
OUT = '/home/user/verbose-guide/deck_content.json'
# slide index -> (ticker, language, shape id carrying the company name)
SLIDES = [(0, 'STB', 'EN', 3), (1, 'STB', 'VN', 2), (2, 'FPT', 'EN', 2), (3, 'FPT', 'VN', 5)]


def table(sh):
    return [[c.text.strip() for c in r.cells] for r in sh.table.rows]


def main():
    prs = Presentation(DECK)
    out = []
    for idx, ticker, lang, name_id in SLIDES:
        sh = {s.shape_id: s for s in prs.slides[idx].shapes}
        points = [p.text.strip() for p in sh[15].text_frame.paragraphs if p.text.strip()]
        out.append({
            'ticker': ticker, 'lang': lang,
            'company': sh[name_id].text_frame.text.strip(),
            'sector': sh[4].text_frame.text.strip(),
            'headline': sh[14].text_frame.text.strip(),
            'analyst': sh[6 if ticker == 'FPT' else 17].text_frame.text.strip(),
            'rating': table(sh[13]),
            'keydata': table(sh[12]),
            'perf': table(sh[11]),
            'points_title': points[0],
            'points': points[1:],
            'fy': table(sh[16]),
        })
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    for s in out:
        print('%s %s  %-52s  %d paragraphs, FY table %dx%d'
              % (s['ticker'], s['lang'], s['company'][:52], len(s['points']),
                 len(s['fy']), len(s['fy'][0])))


if __name__ == '__main__':
    main()
