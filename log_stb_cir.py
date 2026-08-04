# -*- coding: utf-8 -*-
"""Append the 5.5%-NPL / 40-38-36% CIR revision to the internal audit trail."""
import openpyxl
from openpyxl.styles import Font
import update_aug_deck as U

WB = '/home/user/verbose-guide/INTERNAL_Sources_and_Assumptions_Aug26.xlsx'

ROWS = [
    ('', 'STB + FPT — REVISION 04/08/2026 (bản deck August2026)', '', '', '', '', ''),
    ('A1', 'Nguồn số liệu STB', 'model đã tính lại', '', 'MODEL',
     '8ce85368-FinModel_STB_2Q26.xlsx do CV mở bằng Excel',
     'Mọi số STB trên slide đọc thẳng từ file này (update_aug_deck.py), không gõ lại'),
    ('A2', 'CV hạ tỷ lệ xóa nợ FY27F / FY28F', '-1.2% / -0.7%', '% dư nợ', 'MAS',
     'Model!Z287 (từ -0.9%), AA287 (từ -0.5%)',
     'Dự phòng FY27F 8,229 -> 10,595 và FY28F 5,816 -> 7,628. LNTT FY27F 14,676 -> 12,293, '
     'FY28F 21,145 -> 20,064. Tỷ lệ trích/xóa vẫn 1.0x nên dự phòng đã trích không đổi'),
    ('A3', 'LNTT / LNST FY26F',
     '%.0f / %.0f' % (U.PBT[0], U.NPATMI[0]), 'VNDbn', 'MODEL',
     'Model!Y144 / Y153',
     '%+.1f%% CK. Mục tiêu đặt là 8,150; Excel ra %.0f — lệch %.0f tỷ do vòng phản hồi lợi '
     'nhuận -> vốn chủ -> tài sản -> thu nhập lãi. Lấy số của model'
     % (U.PBT_YOY, U.PBT[0], 8150 - U.PBT[0])),
    ('A4', 'CIR FY26F / FY27F / FY28F',
     '%.1f / %.1f / %.1f' % U.CIR, '%', 'MODEL', 'Model!Y135 / (Y136+Y135)',
     'FY26F và FY27F đúng mục tiêu 40/38%%. FY28F ra %.1f%% thay vì 36%% do TOI tăng '
     '(42,858 vs 42,126 khi giải hệ số). CHƯA GIẢI LẠI — chờ CV quyết' % U.CIR[2]),
    ('A5', 'Giá mục tiêu STB', '81,400', 'VND', 'MAS',
     'Ô khuyến nghị slide 0/1 do CV đặt (từ 77,800)',
     'Thị giá 74,100 (04/08/26) -> lợi nhuận kỳ vọng 10%. P/E và P/B toàn bộ bảng FY tính '
     'lại trên giá này, gồm cả các năm quá khứ vì hàng này vốn dựng theo giá mục tiêu'),
    ('A6', 'Bao phủ FY26F', '%.1f%%' % U.COV26, '%', 'DERIVED',
     'Model!Y286 %.0f / NPL 37,941' % U.RESERVE_26,
     'Trích thêm 2H26 %.1f nghìn tỷ (dự phòng cả năm %.0f - 1H26 7,119)'
     % (U.H2_PROV / 1000, U.PROV[0])),
    ('A7', 'Tăng trưởng EPS 26F trên ô tóm tắt', '%.1f' % U.EPS_GROWTH, '%', 'DERIVED',
     'EPS 26F %.0f / EPS 25 2,883' % U.EPS[0],
     'ĐÃ SỬA từ 10.8 mà CV điền — con số đó không khớp bất kỳ dòng nào của model'),
    ('F1', 'FPT: hàng P/E đưa hết về giá mục tiêu', '18.9/18.0/17.3', 'x', 'MAS',
     'Giá mục tiêu 87,950 chia EPS từng năm',
     'Trước đây FY23-25 tính theo giá từng năm (19.7/18.8/18.1) còn FY26F-28F đã theo giá '
     'mục tiêu — hàng này lẫn hai cách. Nay thống nhất, giống cách dựng của STB'),
    ('F2', 'FPT: hàng P/B vẫn lẫn hai cách', '4.7/4.5/4.3', 'x', 'GAP',
     'FY23-25 theo giá từng năm; FY26F-28F theo giá mục tiêu (87,950/25,116 = 3.5)',
     'CHƯA SỬA — CV chỉ yêu cầu P/E. Lỗi hoàn toàn tương tự, sửa hay không tùy CV'),
    ('G20', 'Model!AA117 thiếu dòng NII', '', '', 'GAP',
     'AA117 = +AA124+AA131 trong khi Y117/Z117 = Y121+Y124+Y131',
     'Lỗi có sẵn: ô TOI cột FY28F bỏ mất thu nhập lãi thuần, ra 8,779 thay vì 42,858. '
     'Chỉ là ô hiển thị, không có gì phía sau đọc nó, nhưng nên sửa'),
    ('C3', 'Kế hoạch LNTT FY26 8,100 tỷ', 'suy ra', 'VNDbn', 'CONFLICT',
     'Suy ra từ 4,136/51% theo trích dẫn báo chí "hoàn thành 51% mục tiêu"',
     'CHƯA CÓ NGHỊ QUYẾT ĐHĐCĐ TRONG HỒ SƠ. Slide đang neo dự phóng vào con số này'),
    ('C4', 'Slide ghi "ban lãnh đạo dự kiến ~5.6%" nợ xấu', '5.6%', '%', 'CONFLICT',
     'Dòng 10 ghi 5.6% xuất hiện ở bản 24/07 dưới dạng "ước ~5.6%" — ước tính của MAS',
     'CHƯA SỬA — nếu không có công bố của STB thì phải viết lại'),
    ('G15', 'Giá mục tiêu STB vs sheet Valuation', '81,400 vs ~65,600', 'VND', 'CONFLICT',
     'Valuation!B1 = P/B hợp lý FY27F x BPS FY27F',
     'Khoảng cách nới rộng thêm khi giá mục tiêu lên 81,400 — CHƯA ĐỊNH GIÁ LẠI'),
    ('G16', 'Giá mục tiêu ghi 77,500 trong file stock pick', '77,500', 'VND', 'CONFLICT',
     'Stock Pick!D6 và Target Price!C13 vẫn ghi 77,500',
     'CHƯA SỬA — nay lệch 3,900đ so với 81,400 trên slide, cần CV xác nhận'),
    ('G17', 'Tăng trưởng tín dụng FY26F vs 1.5% thực hiện 1H26', '', '', 'GAP',
     'Model!Y206 = 10.1%; slide ghi 11.7%', 'CHƯA XỬ LÝ'),
    ('G19', 'NPL hàng 19-21 lệch một dòng ở khối theo quý', '', '', 'GAP',
     'DF19/DG19 = nợ nhóm 2 / nợ xấu trong khi nhãn ghi "3) Substandard"',
     'CHƯA SỬA — lỗi ở mọi cột quý. Không ảnh hưởng tỷ lệ NPL hàng 23'),
]


def main():
    wb = openpyxl.load_workbook(WB)
    ws = wb['STB FPT 31Jul']
    r = ws.max_row + 2
    for vals in ROWS:
        for ci, v in enumerate(vals, start=1):
            if v != '':
                ws.cell(row=r, column=ci, value=v)
        if vals[0] == '':
            ws.cell(row=r, column=2).font = Font(bold=True)
        r += 1
    wb.save(WB)
    print('appended %d rows to "STB FPT 31Jul", now %d rows' % (len(ROWS), ws.max_row))


if __name__ == '__main__':
    main()
