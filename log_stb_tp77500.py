# -*- coding: utf-8 -*-
"""Append the VND77,500 target and the VND8,180bn FY26F PBT to the audit trail."""
import openpyxl
from openpyxl.styles import Font
import update_stb_loans as U
from valuation import V

WB = '/home/user/verbose-guide/INTERNAL_Sources_and_Assumptions_Aug26.xlsx'

ROWS = [
    ('', 'STB — REVISION 06/08/2026 d (giá mục tiêu 77,500; LNTT FY26F 8,180)',
     '', '', '', '', ''),
    ('A30', 'LNTT FY26F', '%.0f' % U.PBT[0], 'VNDbn', 'MAS',
     'Model!Y278 = -Y279*0.81355 (trước là 0.8653); Model!Y144',
     'Từ 7,573 lên %.0f (%+.1f%% CK), lần đầu VƯỢT kế hoạch 8,100 tỷ — cao hơn %.1f%%. '
     'Đòn bẩy vẫn là tỷ lệ trích trên xóa nợ: xóa nợ giữ 11,568 tỷ, phần trích cụ thể '
     'giảm từ 10,010 xuống %.0f tỷ. Chi phí dự phòng cả năm %.0f tỷ (%+.1f%% CK)'
     % (U.PBT[0], U.PBT_YOY, U.VS_PLAN, U.F['specific'][0], U.PROV[0], U.PROV_YOY)),
    ('A31', 'Giá mục tiêu', '%.0f' % U.TP, 'VND', 'MAS',
     'Ô khuyến nghị slide 0/1; Valuation!D74 giải lại về đúng mức này',
     'Từ 75,000. Thị giá 74,100 (04/08/26) nên lợi nhuận kỳ vọng %.1f%%. P/E và P/B '
     'toàn bộ bảng FY tính lại trên giá này, gồm cả các năm quá khứ'
     % (U.TP / U.PRICE * 100 - 100)),
    ('A32', 'SỬA: ô khuyến nghị vẫn ghi 81,400 và lợi nhuận kỳ vọng 10%',
     '77,500 / %.1f%%' % (U.TP / U.PRICE * 100 - 100), 'VND, %', 'MODEL',
     'Bảng shape 13 dòng 1 và dòng 3 trên slide 0 và 1',
     'Ô này chưa từng được cập nhật qua ba lần đổi giá mục tiêu (81,400 -> 77,800 -> '
     '75,000 -> 77,500) và vẫn ghi 81,400 với lợi nhuận kỳ vọng 10%. Nay ghi thẳng từ '
     'biến TP nên không thể tụt lại nữa'),
    ('A33', 'Beta định giá %.3f' % V['beta'], '%.3f' % V['beta'], '', 'MAS',
     'Valuation!B11; chi phí vốn CSH %.3f%%; D74 = %.0f'
     % (V['coe'] * 100, V['fair_value']),
     'Giải ngược từ giá mục tiêu 77,500. Beta đã đi theo dự phóng: 1.02 khi mục tiêu '
     '75,000 trên nền tín dụng 2.3%%, 1.222 khi tín dụng lên 8%%, nay 1.173 khi mục tiêu '
     '77,500 và LNTT FY26F 8,180. Chân P/B %.0f, chân RI %.0f'
     % (V['pb_leg'], V['ri_leg'])),
    ('A34', 'Hệ số chi phí giải lại giữ CIR', '%.1f / %.1f / %.1f%%' % U.CIR, '%', 'MODEL',
     'Model!Y/Z/AA132,134 = 1.0301 / 1.0929 / 1.0771',
     'Giảm trích lập làm dự phòng đã trích thấp đi, dư nợ ròng cao lên, thu nhập lãi '
     'nhích lên nên TOI đổi — phải giải lại hệ số để CIR vẫn đúng 42.1/40/38%'),
    ('A35', 'SỬA: vốn chủ FY26F chưa chạy theo lợi nhuận', '%.0f' % U.EQUITY[0],
     'VNDbn', 'MODEL', 'model_read: vốn chủ FY26F = Y85 - Y153 + LNST mới',
     'Trước đây vốn chủ FY26F lấy thẳng cache Y85, đúng khi LNST FY26F còn khớp cache. '
     'Từ khi tín dụng và trích lập đổi thì không còn đúng. Nay dựng lại từ nền: '
     'Y87 + Y94 + (Y98 - Y153) + Y101, cộng LNST của chính bản dự phóng'),
    ('G28', 'Bao phủ nợ xấu FY26F còn %.1f%%' % U.COV26, '%.1f%%' % U.COV26, '%', 'GAP',
     'Dự phòng %.0f trên nợ xấu %.0f' % (U.RESERVE26, U.NPL26),
     'Đây là cái giá của LNTT 8,180: phần trích thêm cắt đi chảy thẳng vào lợi nhuận '
     'nên bộ đệm mỏng đi. 50.9%% -> 48.6%% (khi tín dụng lên 8%%) -> %.1f%%. Bao phủ '
     'FY28F cũng lùi từ 104.6%% về %.1f%%. Nếu CV muốn giữ bao phủ 50%% thì LNTT FY26F '
     'phải lùi về khoảng 7,300 tỷ' % (U.COV26, U.F['coverage'][2])),
]


def main():
    wb = openpyxl.load_workbook(WB)
    ws = wb['STB FPT 31Jul']
    r = ws.max_row + 2
    for vals in ROWS:
        for ci, v in enumerate(vals, start=1):
            if v != '':
                ws.cell(row=r, column=ci, value=v)
        if vals[0] == '':
            ws.cell(row=r, column=2).font = Font(bold=True)
        r += 1
    wb.save(WB)
    print('appended %d rows to "STB FPT 31Jul", now %d rows' % (len(ROWS), ws.max_row))


if __name__ == '__main__':
    main()
