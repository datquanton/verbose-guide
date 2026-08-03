# -*- coding: utf-8 -*-
"""Trim each slide's investment-points block toward ~470 words."""
import shutil, sys
from pptx import Presentation
DECK='/home/user/verbose-guide/MASVN_RS_WM_2H26_outlook_Equity_VN_2026_STBFPT_31July2026.pptx'

FPT_EN={1:("1H26 results tracking ahead of plan: revenue reached VND26,269bn (+12.6% YoY) and PBT "
 "VND5,714bn (+18.1% YoY), completing 45%/49% of the FY26 revenue/PBT plan, with the group PBT "
 "margin widening to 21.8%. In 2Q26, revenue was VND13,789bn, PBT VND2,910bn and NPATMI "
 "VND2,568bn (+14% YoY). Technology revenue reached VND23,138bn (+15% YoY; 88% of group) and PBT "
 "VND3,314bn (+16.9% YoY) on a 14.3% margin; Education, investment and others fell 2.1% to "
 "VND3,131bn yet contributed 42% of group PBT on a ~77% margin."),
 2:("Public-sector digitalization underpins the domestic pipeline: Domestic IT was the "
 "fastest-growing segment in 6M26 at VND4,236bn (+22.5% YoY), with segment PBT doubling to "
 "VND308bn on wins under the national digital-transformation programme — multi-year work with "
 "limited exposure to the global IT cycle, though its 7.3% margin dilutes the blend. Signed "
 "contract value for Global IT reached VND26,338bn in 1H26, +32.3% YoY, including 14 deals above "
 "USD10mn — a book-to-bill of 1.39x against 1.07-1.22x through FY21-25. On that basis we look for "
 "Global IT to reaccelerate to +16.7% in FY27F and +17.7% in FY28F, from +14.7% in FY26F and "
 "+13.4% delivered in 1H26, with the US book cut to +5.9% on price competition."),
 4:("Projection and valuation: we forecast FY26F NPATMI of VND10,944bn (+15.6% YoY), FY27F "
 "VND12,674bn (+15.8%) and FY28F VND14,774bn (+16.6%). Associate income — FPT Telecom (45.66%), "
 "equity-accounted from FY26, with FPT Retail, Synnex FPT and FPT Online — rises to VND2,608bn in "
 "FY26F and VND3,449bn by FY28F, and grew 40.8% YoY in 1H26. Operating leverage carries the rest: "
 "SG&A falls from 17.6% of sales to 16.5% by FY28F, with gross margin unchanged. ROE holds near "
 "27% on net cash of around 60% of equity and a maintained DPS of VND2,000. We raise our DCF "
 "(FCFF) target price to VND87,950 from VND78,750, cutting the assumed cost of debt to 9.0% from "
 "11.5% — FPT's realised funding cost in 1H26 was 4.0% — which lowers WACC to 11.4%. The equity "
 "risk premium stays at 8.97% and long-term growth at 1.5% to carry the risk that AI lowers "
 "prices for IT services. At VND71,700 the stock trades at 11.2x FY26F earnings against a "
 "five-year median above 18x; at the target price, 13.7x FY26F and 11.8x FY27F, implying 23% "
 "upside (BUY). AI newsflow remains supportive: an expanded Microsoft partnership across ASEAN, "
 "Japan and Korea, an AI contract worth tens of millions of USD with a European materials group, "
 "and data-intermediary certification with the Ministry of Public Security's National Data "
 "Centre. Risks: (1) AI-driven price deflation outpacing our long-term growth assumption (2) "
 "subdued US and global IT spending and tariff-related contract delays (3) JPY/USD volatility.")}

FPT_VN={1:("KQKD 1H26 vượt tiến độ kế hoạch: doanh thu đạt 26,269 tỷ đồng (+12.6% CK) và LNTT 5,714 "
 "tỷ đồng (+18.1% CK), hoàn thành 45%/49% kế hoạch năm; biên LNTT tập đoàn mở rộng lên 21.8%. "
 "Riêng Q2/2026, doanh thu 13,789 tỷ đồng, LNTT 2,910 tỷ đồng và LNST-CĐTS 2,568 tỷ đồng (+14% "
 "CK). Mảng Công nghệ đạt 23,138 tỷ đồng (+15% CK; 88% doanh thu tập đoàn) và LNTT 3,314 tỷ đồng "
 "(+16.9% CK), biên 14.3%; Giáo dục, đầu tư và khác giảm 2.1% về 3,131 tỷ đồng nhưng vẫn đóng góp "
 "42% LNTT tập đoàn."),
 2:("Số hóa khu vực công là nền tảng của danh mục trong nước: CNTT trong nước tăng nhanh nhất "
 "trong 6T26, đạt 4,236 tỷ đồng (+22.5% CK), LNTT mảng tăng gấp đôi lên 308 tỷ đồng nhờ trúng "
 "thầu thuộc chương trình chuyển đổi số quốc gia — nguồn việc kéo dài nhiều năm, ít phụ thuộc chu "
 "kỳ CNTT toàn cầu, tuy biên 7.3% làm loãng biên chung. Doanh thu ký mới CNTT nước ngoài đạt "
 "26,338 tỷ đồng trong 1H26 (+32.3% CK), gồm 14 dự án trên 10 triệu USD — tỷ lệ ký mới/doanh thu "
 "1.39 lần so với 1.07-1.22 lần giai đoạn FY21-25. Trên cơ sở đó, chúng tôi dự phóng CNTT nước "
 "ngoài tăng lên +16.7% FY27F và +17.7% FY28F, từ +14.7% FY26F và +13.4% của 1H26; thị trường Mỹ "
 "hạ về +5.9% do cạnh tranh giá thầu."),
 4:("Dự phóng và định giá: chúng tôi dự phóng LNST-CĐTS FY26F đạt 10,944 tỷ đồng (+15.6% CK), "
 "FY27F 12,674 tỷ đồng (+15.8%) và FY28F 14,774 tỷ đồng (+16.6%). Lợi nhuận từ công ty liên kết — "
 "FPT Telecom (45.66%, hạch toán theo phương pháp VCSH từ FY26), FPT Retail, Synnex FPT và FPT "
 "Online — tăng lên 2,608 tỷ đồng FY26F và 3,449 tỷ đồng FY28F, và đã tăng 40.8% CK trong 1H26. "
 "Phần còn lại đến từ đòn bẩy hoạt động: chi phí bán hàng và quản lý giảm từ 17.6% doanh thu "
 "xuống 16.5% vào FY28F, biên gộp giữ nguyên. ROE duy trì quanh 27% nhờ tiền ròng khoảng 60% VCSH "
 "và cổ tức tiền mặt 2,000 đồng/cp. Chúng tôi nâng giá mục tiêu DCF (FCFF) lên 87,950 đồng từ "
 "78,750 đồng, hạ giả định chi phí nợ vay về 9.0% từ 11.5% — chi phí lãi vay thực tế 1H26 là 4.0% "
 "— đưa WACC xuống 11.4%. Phần bù rủi ro vốn cổ phần giữ 8.97% và tăng trưởng dài hạn giữ 1.5% để "
 "phản ánh rủi ro AI làm giảm giá dịch vụ CNTT. Tại 71,700 đồng, cổ phiếu giao dịch ở 11.2x P/E "
 "FY26F so với trung vị 5 năm trên 18x; tại giá mục tiêu là 13.7x FY26F và 11.8x FY27F, hàm ý "
 "tiềm năng tăng 23% (MUA). Rủi ro: (1) giảm phát giá do AI nhanh hơn giả định tăng trưởng dài "
 "hạn (2) chi tiêu CNTT tại Mỹ và toàn cầu thấp cùng hợp đồng bị trì hoãn do thuế quan (3) biến "
 "động tỷ giá JPY/USD.")}

STB_VN={3:("Chất lượng tài sản suy giảm vượt mức ban lãnh đạo dự kiến: tỷ lệ nợ xấu đạt 7.54% cuối "
 "Q2/2026, tăng 1.13%p so với đầu năm và cao nhất hệ thống, so với ước tính khoảng 5.6% của ban "
 "lãnh đạo. Dư nợ nhóm 3-5 tăng khoảng 7,800 tỷ đồng so với đầu năm và 6,500 tỷ đồng so với quý "
 "trước, lên 47,957 tỷ đồng, trong đó khoảng 67% (32,000 tỷ đồng) xếp nhóm 5, nợ nhóm 4 tăng hơn "
 "gấp đôi (+156%). Giả định nợ xấu FY26F dưới 4.5% không còn khả thi nếu không đẩy mạnh xóa nợ. "
 "Chi phí dự phòng FY26F 10.4 nghìn tỷ đồng (-8.9% CK) sát nhịp thực hiện 1H26; điểm yếu hơn nằm "
 "ở FY27F, hàm ý chi phí tín dụng chỉ khoảng 0.8% trong khi nợ xấu vẫn trên 7%. Cả hai đang được "
 "rà soát, cùng khuyến nghị NẮM GIỮ và giá mục tiêu 77,800 đồng (2.4x P/B FY26F)."),
 6:("Các giả định FY26F còn lại giữ nguyên: NII 27,010 tỷ đồng (+1.2% CK) với tăng trưởng tín "
 "dụng 11.7%, NIM 2.94% (-38bps CK), CIR 42.7% (+2.0%p CK) và thu nhập khác 1,327 tỷ đồng (+5% "
 "CK); dự phóng LNTT FY26F vẫn cao hơn kế hoạch 8,100 tỷ đồng đã được ĐHĐCĐ thông qua. Ở diễn "
 "biến khác, Sacombank đã thu giữ 507 giấy chứng nhận quyền sử dụng đất thuộc dự án Viva City "
 "(Đồng Nai) đối với dư nợ gốc quá hạn gần 350 tỷ đồng; quy mô chưa trọng yếu nhưng thể hiện kênh "
 "xử lý tài sản đảm bảo mà đà hồi phục 2H26 phụ thuộc vào.")}

shutil.copy(DECK,DECK+'.bak')
prs=Presentation(DECK)
sid=lambda sl,i: next(s for s in sl.shapes if s.shape_id==i)
for slide_i,spec in [(2,FPT_EN),(3,FPT_VN),(1,STB_VN)]:
    tf=sid(prs.slides[slide_i],15).text_frame
    for pi,txt in spec.items():
        para=tf.paragraphs[pi]
        assert para.runs, (slide_i,pi)
        para.runs[0].text=txt
        for r in para.runs[1:]: r.text=''
prs.save(DECK)

prs=Presentation(DECK)
print('%-10s %8s %8s' % ('slide','words','target'))
for i,nm in enumerate(['STB EN','STB VN','FPT EN','FPT VN']):
    tf=sid(prs.slides[i],15).text_frame
    tot=sum(len(''.join(r.text for r in p.runs).split()) for p in tf.paragraphs)
    print('%-10s %8d %8s   %s' % (nm,tot,'~470','OK' if abs(tot-470)<=45 else 'off by %+d'%(tot-470)))
