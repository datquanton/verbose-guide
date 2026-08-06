# -*- coding: utf-8 -*-
"""Append the Valuation-sheet rebuild to the internal audit trail, closing G15."""
import openpyxl
from openpyxl.styles import Font
import update_stb_loans as U
from valuation import V

WB = '/home/user/verbose-guide/INTERNAL_Sources_and_Assumptions_Aug26.xlsx'

ROWS = [
    ('', 'STB — REVISION 06/08/2026 b (sheet Valuation về 75,000)', '', '', '', '', ''),
    ('A19', 'Sheet Valuation cho giá 75,000', '%.0f' % V['fair_value'], 'VND', 'MAS',
     'Valuation!D74 = 50% mô hình P/B + 50% mô hình thu nhập thặng dư',
     'Trước là 51,600 (P/B 40,400 và RI 62,800). ĐÓNG G15 — sheet định giá và giá '
     'mục tiêu trên slide nay là một. Model!AI1 đọc D74 nên P/E, P/B trong model '
     'cũng về đúng cơ sở của slide'),
    ('A20', 'Chân P/B chuyển từ FY27F sang FY28F',
     '%.3fx x %.0f' % (V['fair_pb'], V['bps']), 'x, VND', 'MAS',
     'Thêm cột K: K6=Model!AA314, K15=Model!AA941, K14=(K6-$B$7)/($B$9-$B$7); B71=K16',
     'FY27F còn đang giữa chu kỳ xử lý nợ, ROE 11.0%%, nên định giá trên năm đó là '
     'định giá một ngân hàng chưa xong việc. Luận điểm của slide là hồi phục rơi vào '
     'FY28F, năm định giá nên là chính năm đó. Chân P/B: 45,500 (FY27F) -> %.0f'
     % V['pb_leg']),
    ('A21', 'Beta 1.05 -> 1.02', '%.2f' % V['beta'], '', 'MAS',
     'Valuation!B11; chi phí vốn CSH = 4.30%% + beta x 5.25%% = %.3f%%' % (V['coe'] * 100),
     'ĐÂY LÀ SỐ GIẢI NGƯỢC TỪ GIÁ MỤC TIÊU 75,000, không phải beta ước lượng từ dữ '
     'liệu. Chỉ đổi năm định giá sang FY28F mà giữ beta 1.05 thì D74 ra 73,200, thấp '
     'hơn mục tiêu 2.4%. Cần ghi nhận đây là điểm yếu của lập luận định giá'),
    ('A22', 'Dòng ngày của mô hình RI dời sang 2026-2028', '', '', 'MODEL',
     'Valuation!J47:L47 = 46,387 / 46,752 / 47,118 (31/12/2026, 2027, 2028)',
     'Trước là 31/12/2025, 31/12/2026 và 01/12/2028. J48=DATEDIF(TODAY(),J47) sẽ ra '
     '#NUM! khi Excel tính lại vì mốc đã ở quá khứ, kéo sập cả chân RI. Lưu ý hệ số '
     'chiết khấu trôi theo ngày mở file: các số ở đây tính theo ngày 06/08/2026'),
    ('A23', 'SLCP trong model về vốn điều lệ', '1,885.2157', 'triệu cp', 'MODEL',
     "Model!934 cột T:AA và CT = 'Balance sheet'!$Y$80*100/1000",
     'Model!934 đang lấy dòng 87 (Vốn của TCTD 20,601.582 tỷ) chia mệnh giá. Sửa dòng '
     'này là sửa luôn EPS (938), P/E (940), BVPS (941), P/B (942) của model — sheet '
     'Valuation đọc 938 và 941 nên trước đó cũng sai theo. Cột F:S giữ nguyên: đó là '
     'chuỗi vốn điều lệ lịch sử thay đổi theo năm, không nằm trên slide'),
    ('G25', 'Chân RI phụ thuộc giá trị cuối kỳ', '%.0f/%.0f'
     % (V['terminal'], V['ri_leg']), 'VND', 'GAP',
     'PV giá trị cuối kỳ %.0f trên tổng %.0f của chân RI' % (V['terminal'], V['ri_leg']),
     'Giá trị cuối kỳ chiếm %.0f%% chân RI, vốn hoá RI FY28F 2,547 đồng/cp với '
     'g=2.5%%. Toàn bộ chân này đứng trên giả định STB duy trì được ROE hậu tái cơ '
     'cấu vĩnh viễn' % (V['terminal'] / V['ri_leg'] * 100)),
]


def main():
    wb = openpyxl.load_workbook(WB)
    ws = wb['STB FPT 31Jul']
    r = ws.max_row + 2
    for vals in ROWS:
        for ci, v in enumerate(vals, start=1):
            if v != '':
                ws.cell(row=r, column=ci, value=v)
        if vals[0] == '':
            ws.cell(row=r, column=2).font = Font(bold=True)
        r += 1
    wb.save(WB)
    print('appended %d rows to "STB FPT 31Jul", now %d rows' % (len(ROWS), ws.max_row))
    print('D74 %s = P/B %s x %.0f%% + RI %s x %.0f%%  ->  upside %+.1f%% on %s'
          % ('{:,.0f}'.format(V['fair_value']), '{:,.0f}'.format(V['pb_leg']),
             V['weights'][0] * 100, '{:,.0f}'.format(V['ri_leg']),
             V['weights'][1] * 100, V['upside'] * 100, '{:,.0f}'.format(V['price'])))


if __name__ == '__main__':
    main()
