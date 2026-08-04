# -*- coding: utf-8 -*-
"""The chart data behind the mobile report, as one workbook.

One series per row, years across the columns, so any block can be selected and
charted in Excel directly.  This is also the single source the native
PowerPoint charts in build_mobile.py are built from - both read SERIES here, so
the deck and the workbook cannot disagree.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

OUT = '/home/user/verbose-guide/Mobile_Report_ChartData_STB_FPT_2Q26.xlsx'
YRS = ['FY24', 'FY25', 'FY26F', 'FY27F', 'FY28F']

# (ticker, chart key, series name EN, series name VN, values)
SERIES = [
    ('STB', 'is', 'Net interest income', 'Thu nhập lãi thuần',
     [24532, 26681, 26704, 29545, 34079]),
    ('STB', 'is', 'Non-interest income', 'Thu nhập ngoài lãi',
     [4145, 5376, 5950, 7382, 8779]),
    ('STB', 'pbt', 'Profit before tax', 'Lợi nhuận trước thuế',
     [12720, 7628, 8143, 12293, 20064]),
    ('STB', 'val_pe', 'P/E at target price (x)', 'P/E theo giá mục tiêu (lần)',
     [15.8, 28.2, 26.5, 17.5, 10.7]),
    ('STB', 'val_pb', 'P/B at target price (x)', 'P/B theo giá mục tiêu (lần)',
     [3.1, 2.8, 2.5, 2.2, 1.8]),
    ('FPT', 'rev', 'Revenue', 'Doanh thu',
     [62849, 70208, 57284, 66048, 76654]),
    ('FPT', 'is', 'Profit before tax', 'Lợi nhuận trước thuế',
     [11070, 13134, 13070, 15159, 17671]),
    ('FPT', 'is', 'NPATMI', 'LNST-CĐTS',
     [7857, 9464, 10944, 12674, 14774]),
    ('FPT', 'val_pe', 'P/E at target price (x)', 'P/E theo giá mục tiêu (lần)',
     [18.0, 17.3, 13.7, 11.8, 10.1]),
    ('FPT', 'val_pb', 'P/B at target price (x)', 'P/B theo giá mục tiêu (lần)',
     [4.5, 4.3, 3.5, 3.0, 2.5]),
]

NAVY, AMBER = '01437C', 'F38120'
NOTES = [
    'Units: VNDbn except P/E and P/B, which are multiples.',
    'STB is on the recalculated FinModel_STB_2Q26 (Model!Y121, Y124+Y131, Y144).',
    'FPT is on FPT_2Q26, Report sheet rows 3, 15 and 20.',
    'FPT FY26F revenue is not comparable with FY25 as reported: FPT Telecom is '
    'equity-accounted from FY26F. Against a restated FY25 of 50,607 the change is +13.2%.',
    'P/E and P/B are struck on the target prices used in the August deck: '
    'STB VND81,400, FPT VND87,950. Historical years use the same target price, '
    'which is how the deck FY tables are built.',
]


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Chart data'
    ws['A1'] = 'Mobile report — STB & FPT 2Q26: chart data'
    ws['A1'].font = Font(bold=True, size=13, color=NAVY)
    ws['A2'] = 'Select a block including its header row and insert a column chart.'
    ws['A2'].font = Font(italic=True, size=9, color='595959')

    hdr = ['Ticker', 'Chart', 'Series (EN)', 'Series (VN)'] + YRS
    r = 4
    for ci, h in enumerate(hdr, start=1):
        c = ws.cell(row=r, column=ci, value=h)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(horizontal='center')
    for tick, key, en, vn, vals in SERIES:
        r += 1
        ws.cell(row=r, column=1, value=tick).font = Font(bold=True, color=AMBER)
        ws.cell(row=r, column=2, value=key)
        ws.cell(row=r, column=3, value=en)
        ws.cell(row=r, column=4, value=vn)
        for ci, v in enumerate(vals, start=5):
            c = ws.cell(row=r, column=ci, value=v)
            c.number_format = '#,##0.0' if key.startswith('val') else '#,##0'
    r += 2
    ws.cell(row=r, column=1, value='Notes').font = Font(bold=True, color=NAVY)
    for n in NOTES:
        r += 1
        ws.cell(row=r, column=1, value=n).font = Font(size=9, color='595959')

    for ci, w in enumerate([9, 9, 26, 30] + [11] * len(YRS), start=1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = 'E5'
    wb.save(OUT)
    print('wrote %s  (%d series)' % (OUT, len(SERIES)))


if __name__ == '__main__':
    main()
